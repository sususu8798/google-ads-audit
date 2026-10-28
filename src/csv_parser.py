"""Google Ads 报表解析与 URL 聚合。

设计要点（PRD 第 4.3 / 4.4 / 4.8 节）：
- 自动跳过前置元数据行（报表名称、日期范围），定位真正字段表头
- 同义字段映射；无法识别的列保留原列名
- 数值稳健解析：百分比、货币、千分位、破折号、空值、负值
- 按正确数据粒度聚合；不直接平均行级百分比
- 多币种分别汇总，禁止合并；零分母显示 N/A
- 检测重复行 / 汇总行，避免重复计算
"""
import csv
import io
import re
from collections import defaultdict
from typing import List, Optional

from urllib.parse import urlsplit, urlunsplit

EXPECTED_FIELDS = [
    "final_url", "mobile_final_url", "campaign", "ad_group", "ad_status",
    "clicks", "impressions", "ctr", "currency", "avg_cpc", "cost",
    "conv_rate", "conversions", "cost_per_conv", "status_reasons", "ad_type",
]

# 列名（规范化后）-> 标准字段
HEADER_MAP = {
    "final url": "final_url",
    "mobile final url": "mobile_final_url",
    "campaign": "campaign",
    "campaign name": "campaign",
    "ad group": "ad_group",
    "ad group name": "ad_group",
    "ad status": "ad_status",
    "status": "ad_status",
    "clicks": "clicks",
    "impr.": "impressions",
    "impressions": "impressions",
    "ctr": "ctr",
    "currency code": "currency",
    "currency": "currency",
    "avg. cpc": "avg_cpc",
    "average cpc": "avg_cpc",
    "cost": "cost",
    "conv. rate": "conv_rate",
    "conversion rate": "conv_rate",
    "conversions": "conversions",
    "cost / conv.": "cost_per_conv",
    "cost / conversion": "cost_per_conv",
    "cost per conv.": "cost_per_conv",
    "status reasons": "status_reasons",
    "ad type": "ad_type",
    "type": "ad_type",
}

HEADLINE_RE = re.compile(r"^(?:headline|标题)\s*\d+$", re.I)
DESC_RE = re.compile(r"^(?:description|说明|描述)\s*\d+$", re.I)
DATE_RE = re.compile(r"(\d{4}[-/]\d{1,2}[-/]\d{1,2})")
RANGE_RE = re.compile(r"(\d{4}[-/]\d{1,2}[-/]\d{1,2})\s*(?:至|to|-|–|—)\s*(\d{4}[-/]\d{1,2}[-/]\d{1,2})", re.I)

TRACKING_PARAMS = {"gclid", "fbclid", "utm_source", "utm_medium", "utm_campaign",
                   "utm_term", "utm_content", "utm_id", "gbraid", "wbraid", "msclkid"}


def _norm_header(h: str) -> str:
    return re.sub(r"\s+", " ", (h or "").strip().lower())


def _parse_number(value) -> Optional[float]:
    if value is None:
        return None
    s = str(value).strip()
    if s in ("", "-", "—", "–", "N/A", "n/a", "#N/A", "—"):
        return None
    is_pct = s.endswith("%")
    s = s.replace("%", "")
    # 去除货币符号/字母与千分位
    s = re.sub(r"[^\d.,\-]", "", s)
    s = s.replace(",", "")
    if s in ("", "-", "."):
        return None
    try:
        num = float(s)
    except ValueError:
        return None
    if is_pct:
        return num  # 百分比原值
    return num


def _parse_currency_symbol(value) -> Optional[str]:
    if value is None:
        return None
    s = str(value).strip()
    m = re.match(r"^([A-Za-z]{1,5})\s*[\d.,\-]", s)
    if m:
        return m.group(1).upper()
    # 仅符号如 ¥ $ €
    sym_map = {"¥": "CNY", "$": "USD", "€": "EUR", "£": "GBP"}
    if s and s[0] in sym_map:
        return sym_map[s[0]]
    return None


def _parse_currency_code(value) -> Optional[str]:
    """Currency code 列通常是纯 3 字母代码（如 CNY/USD），无数字。"""
    if value is None:
        return None
    s = str(value).strip()
    if re.fullmatch(r"[A-Za-z]{3,5}", s):
        return s.upper()
    return None


def _normalize_url(url: str) -> Optional[str]:
    if not url:
        return None
    s = str(url).strip()
    if s in ("-", "", "—"):
        return None
    try:
        parts = urlsplit(s)
    except ValueError:
        return None
    if not parts.scheme:
        s = "https://" + s
        parts = urlsplit(s)
    host = parts.netloc.lower().replace("www.", "")
    # 去除 tracking 参数
    from urllib.parse import parse_qsl, urlencode
    items = parse_qsl(parts.query, keep_blank_values=True)
    kept = [(k, v) for k, v in items if k.lower() not in TRACKING_PARAMS]
    query = urlencode(kept)
    path = parts.path or "/"
    return urlunsplit((parts.scheme, host, path, query, ""))


def _read_rows(filename: str, content: bytes):
    """返回 list[list[str]]。自动识别 CSV/XLSX。"""
    lower = filename.lower()
    if lower.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append([("" if c is None else str(c)) for c in r])
        return rows
    # CSV
    text = content.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text.split("\n", 30).__str__()[:4096], delimiters=",;\t|")
        delim = dialect.delimiter
    except Exception:
        delim = ","
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    return [row for row in reader]


def _find_header(rows: List[List[str]]):
    best_idx, best_score = 0, 0
    for i, row in enumerate(rows):
        norm = [_norm_header(c) for c in row]
        score = sum(1 for c in norm if c in HEADER_MAP)
        # 强约束：必须包含 final url 或 campaign 才可能是表头
        if ("final url" in norm or "campaign" in norm) and score > best_score:
            best_score, best_idx = score, i
    return best_idx, best_score


def parse_report_file(filename: str, content: bytes) -> dict:
    rows = _read_rows(filename, content)
    if not rows:
        return {"error": "文件为空或无法解析", "preview": {}, "records": [], "aggregated": [], "warnings": []}

    header_idx, header_score = _find_header(rows)
    header = [_norm_header(c) for c in rows[header_idx]]
    col_map = {}
    for ci, h in enumerate(header):
        if h in HEADER_MAP:
            col_map[ci] = HEADER_MAP[h]

    meta_rows = rows[:header_idx]
    data_rows = rows[header_idx + 1:]

    # 元数据：报表名称 / 日期范围
    report_name = ""
    date_range = None
    for mr in meta_rows:
        line = " ".join(str(c) for c in mr if c).strip()
        if not line:
            continue
        if not report_name and len(line) < 120:
            report_name = line
        m = RANGE_RE.search(line)
        if m:
            date_range = [m.group(1), m.group(2)]

    # 映射标准字段 -> 列索引（含 headline/description 收集）
    # 注意：标题/描述列不在 HEADER_MAP 中，必须扫描所有表头列，否则会漏读
    field_cols = defaultdict(list)
    headline_cols = []
    desc_cols = []
    for ci, h in enumerate(header):
        if HEADLINE_RE.match(h):
            headline_cols.append(ci)
        elif DESC_RE.match(h):
            desc_cols.append(ci)
        elif h in HEADER_MAP:
            field_cols[HEADER_MAP[h]].append(ci)

    detected = sorted(set(col_map.values()))
    missing = [f for f in EXPECTED_FIELDS if f not in detected]
    raw_headers = header  # 保留原始列名用于展示

    # 解析数据行
    records = []
    seen = set()
    dup_count = 0
    dup_effective = 0
    currency_set = set()
    campaigns = set()
    ad_groups = set()
    warnings = []

    summary_row_markers = {"total", "总计", "合计", "summary", "小计", "grand total"}

    for ri, row in enumerate(data_rows):
        if not any(str(c).strip() for c in row):
            continue  # 空行
        # 判定汇总行
        camp_val = ""
        for ci in field_cols.get("campaign", []):
            if ci < len(row):
                camp_val = str(row[ci]).strip()
        is_summary = camp_val.lower() in summary_row_markers
        # final url
        fu_raw = ""
        for ci in field_cols.get("final_url", []):
            if ci < len(row):
                fu_raw = str(row[ci]).strip()
        norm_url = _normalize_url(fu_raw)

        # 去重（基于关键指标行）；短行（列数不足）用空串兜底，防止越界崩溃
        key = (norm_url, camp_val) + tuple(
            (str(row[ci]).strip() if ci < len(row) else "") for ci in
            sorted(field_cols.get("clicks", []) + field_cols.get("impressions", []) +
                   field_cols.get("cost", []) + field_cols.get("conversions", [])))
        if key in seen:
            dup_count += 1
            # 只有真正参与计算的数据行（有 URL、非汇总行）重复才值得提示；
            # 总计行 / 无 URL 行的重复对结果零影响，静默处理
            if norm_url and not is_summary:
                dup_effective += 1
            continue
        seen.add(key)

        def getf(std, idx=0):
            cols = field_cols.get(std, [])
            if idx < len(cols) and cols[idx] < len(row):
                return row[cols[idx]]
            return ""

        rec = {
            "row_index": ri + header_idx + 2,
            "is_summary": is_summary,
            "final_url": norm_url,
            "final_url_raw": fu_raw,
            "mobile_final_url": _normalize_url(getf("mobile_final_url")),
            "campaign": camp_val,
            "ad_group": getf("ad_group"),
            "ad_status": getf("ad_status"),
            "clicks": _parse_number(getf("clicks")),
            "impressions": _parse_number(getf("impressions")),
            "ctr_original": _parse_number(getf("ctr")),
            "avg_cpc": _parse_number(getf("avg_cpc")),
            "cost": _parse_number(getf("cost")),
            "conv_rate": _parse_number(getf("conv_rate")),
            "conversions": _parse_number(getf("conversions")),
            "cost_per_conv": _parse_number(getf("cost_per_conv")),
            "status_reasons": getf("status_reasons"),
            "ad_type": getf("ad_type"),
            "currency": _parse_currency_code(getf("currency")) or _parse_currency_symbol(getf("cost")) or _parse_currency_symbol(getf("currency")) or "UNKNOWN",
            "headlines": [str(row[ci]).strip() for ci in headline_cols if ci < len(row) and str(row[ci]).strip()],
            "descriptions": [str(row[ci]).strip() for ci in desc_cols if ci < len(row) and str(row[ci]).strip()],
        }
        currency_set.add(rec["currency"])
        if rec["is_summary"]:
            continue  # 汇总行不计入维度集合与计数
        if camp_val:
            campaigns.add(camp_val)
        if rec["ad_group"]:
            ad_groups.add(rec["ad_group"])
        records.append(rec)

    if dup_effective:
        warnings.append(f"检测到 {dup_effective} 行重复记录，已去重避免重复计算")
    if missing:
        warnings.append(f"缺少关键字段：{', '.join(missing)}（将影响部分分析）")
    if "UNKNOWN" in currency_set:
        warnings.append("部分行未识别币种，已标记为 UNKNOWN；多币种请确认汇率")

    # 聚合（按 URL + 币种），保留追溯
    agg = defaultdict(lambda: {
        "clicks": 0.0, "impressions": 0.0, "cost": 0.0, "conversions": 0.0,
        "campaigns": set(), "ad_groups": set(), "headlines": [], "descriptions": [],
        "ad_statuses": set(), "rows": 0,
    })
    no_url_rows = 0
    for r in records:
        if r["is_summary"]:
            continue
        if not r["final_url"]:
            no_url_rows += 1
            continue
        key = (r["final_url"], r["currency"])
        a = agg[key]
        a["clicks"] += r["clicks"] or 0
        a["impressions"] += r["impressions"] or 0
        a["cost"] += r["cost"] or 0
        a["conversions"] += r["conversions"] or 0
        if r["campaign"]:
            a["campaigns"].add(r["campaign"])
        if r["ad_group"]:
            a["ad_groups"].add(r["ad_group"])
        a["headlines"].extend(r["headlines"])
        a["descriptions"].extend(r["descriptions"])
        if r["ad_status"]:
            a["ad_statuses"].add(r["ad_status"])
        a["rows"] += 1

    aggregated = []
    for (url, currency), a in agg.items():
        ctr = round(a["clicks"] / a["impressions"], 4) if a["impressions"] else None
        avg_cpc = round(a["cost"] / a["clicks"], 2) if a["clicks"] else None
        cpa = round(a["cost"] / a["conversions"], 2) if a["conversions"] else None  # 零转化 -> None (N/A)
        conv_rate = round(a["conversions"] / a["clicks"], 4) if a["clicks"] else None
        aggregated.append({
            "final_url": url,
            "currency": currency,
            "clicks": a["clicks"],
            "impressions": a["impressions"],
            "cost": a["cost"],
            "conversions": a["conversions"],
            "ctr": ctr,
            "avg_cpc": avg_cpc,
            "cpa": cpa,
            "conv_rate": conv_rate,
            "campaign_count": len(a["campaigns"]),
            "ad_group_count": len(a["ad_groups"]),
            "source_rows": a["rows"],
            "campaigns": sorted(a["campaigns"]),
            "ad_groups": sorted(a["ad_groups"]),
            "headlines_sample": a["headlines"][:15],
            "descriptions_sample": a["descriptions"][:6],
            "ad_statuses": sorted(a["ad_statuses"]),
        })

    preview = {
        "filename": filename,
        "report_name": report_name,
        "date_range": date_range,
        "data_rows": len([r for r in records if not r["is_summary"]]),
        "column_count": len(header),
        "url_count": len({k[0] for k in agg.keys()}),
        "campaign_count": len(campaigns),
        "ad_group_count": len(ad_groups),
        "currencies": sorted(currency_set),
        "detected_fields": detected,
        "missing_fields": missing,
        "duplicate_count": dup_count,
        "rows_without_url": no_url_rows,
        "header_row_index": header_idx,
        "raw_headers": [h for h in raw_headers if h],
    }

    return {
        "preview": preview,
        "records": records,
        "aggregated": aggregated,
        "warnings": warnings,
        "errors": [],
    }
